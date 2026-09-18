import io
import os
import httpx
from typing import Optional
import pypdf
import docx
from ai_service.core.logging import logger


class DocumentParser:
    """
    High-fidelity document parser supporting PDF, DOCX, and plain text resumes.
    Supports parsing from in-memory bytes, local files, or remote URLs.
    """

    @staticmethod
    def parse_from_bytes(content: bytes, file_type_or_name: str) -> str:
        """
        Extract text from raw byte content based on extension or mime type.
        """
        lower = file_type_or_name.lower()

        if lower.endswith(".pdf") or "pdf" in lower:
            return DocumentParser._parse_pdf_bytes(content)
        elif lower.endswith(".docx") or "docx" in lower or "wordprocessing" in lower:
            return DocumentParser._parse_docx_bytes(content)
        else:
            # Fallback to plain text decoding
            try:
                return content.decode("utf-8")
            except UnicodeDecodeError:
                return content.decode("latin-1", errors="ignore")

    @staticmethod
    def parse_from_file(file_path: str) -> str:
        """
        Extract text from a file stored on the local disk.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Resume document not found at: {file_path}")

        with open(file_path, "rb") as f:
            content = f.read()

        return DocumentParser.parse_from_bytes(content, file_path)

    @staticmethod
    async def parse_from_url(url: str, timeout: float = 15.0) -> str:
        """
        Fetch a document from a remote URL (e.g. GCS, S3, or HTTP) and extract its text content.
        """
        logger.info(f"Fetching document from URL: {url}")
        try:
            async with httpx.AsyncClient(
                timeout=timeout, follow_redirects=True
            ) as client:
                res = await client.get(url)
                if res.status_code != 200:
                    logger.warning(
                        f"Failed to fetch document from URL {url} (HTTP {res.status_code})"
                    )
                    return (
                        f"Document content unavailable. HTTP status: {res.status_code}"
                    )

                content_type = res.headers.get("content-type", "")
                parsed_text = DocumentParser.parse_from_bytes(
                    res.content, url + " " + content_type
                )

                if not parsed_text or not parsed_text.strip():
                    return f"Document at {url} contained no extractable text."

                return parsed_text.strip()
        except Exception as ex:
            logger.error(f"Error fetching/parsing document from URL {url}: {ex}")
            return f"Failed to download and parse resume from URL: {str(ex)}"

    @staticmethod
    async def parse_from_relative_path_or_url(
        path_or_url: str, timeout: float = 15.0
    ) -> str:
        """
        Resolves a file path or URL. First checks local disk candidate paths,
        and falls back to DOTNET_API_BASE_URL HTTP download.
        """
        if path_or_url.startswith("http://") or path_or_url.startswith("https://"):
            return await DocumentParser.parse_from_url(path_or_url, timeout=timeout)

        clean_path = path_or_url.lstrip("/\\")

        # 1. Search candidate paths on local disk
        current_dir = os.path.dirname(os.path.abspath(__file__))
        candidates = [
            path_or_url,
            clean_path,
            os.path.join(os.getcwd(), clean_path),
            os.path.join(os.getcwd(), "..", "api", "HireWise.Api", clean_path),
            os.path.join(os.getcwd(), "apps", "api", "HireWise.Api", clean_path),
            os.path.abspath(
                os.path.join(
                    current_dir,
                    "..",
                    "..",
                    "..",
                    "..",
                    "api",
                    "HireWise.Api",
                    clean_path,
                )
            ),
            os.path.join(
                "c:\\nadil-dulnidu\\HireWise\\apps\\api\\HireWise.Api", clean_path
            ),
        ]

        for candidate in candidates:
            if candidate and os.path.isfile(candidate):
                logger.info(f"Found local resume file at: {candidate}")
                try:
                    return DocumentParser.parse_from_file(candidate)
                except Exception as ex:
                    logger.warning(f"Failed parsing local file {candidate}: {ex}")

        # 2. Try fetching from DOTNET_API_BASE_URL
        from ai_service.core.config import settings

        if settings.DOTNET_API_BASE_URL:
            full_url = f"{settings.DOTNET_API_BASE_URL.rstrip('/')}/{clean_path}"
            logger.info(f"Attempting download from API: {full_url}")
            return await DocumentParser.parse_from_url(full_url, timeout=timeout)

        return f"Candidate applied with resume reference: {path_or_url}"

    @staticmethod
    def _parse_pdf_bytes(content: bytes) -> str:
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(text)
            return "\n\n".join(pages_text)
        except Exception as ex:
            logger.warning(f"PDF parsing error: {ex}. Attempting fallback.")
            try:
                return content.decode("utf-8", errors="ignore")
            except Exception:
                return ""

    @staticmethod
    def _parse_docx_bytes(content: bytes) -> str:
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        c.text.strip() for c in row.cells if c.text.strip()
                    )
                    if row_text:
                        paragraphs.append(row_text)
            return "\n\n".join(paragraphs)
        except Exception as ex:
            logger.warning(f"DOCX parsing error: {ex}")
            return ""
