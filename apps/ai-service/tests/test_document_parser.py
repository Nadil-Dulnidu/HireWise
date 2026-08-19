import io
import pytest
import pypdf
import docx
from ai_service.utils.document_parser import DocumentParser

def test_parse_plain_text():
    text = "John Doe\nSenior Python Engineer with 6 years experience in FastAPI and PostgreSQL."
    result = DocumentParser.parse_from_bytes(text.encode("utf-8"), "resume.txt")
    assert "Senior Python Engineer" in result
    assert "FastAPI" in result

def test_parse_pdf_bytes():
    # Generate a small in-memory valid PDF using pypdf
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    
    stream = io.BytesIO()
    writer.write(stream)
    pdf_bytes = stream.getvalue()

    result = DocumentParser.parse_from_bytes(pdf_bytes, "candidate_resume.pdf")
    assert isinstance(result, str)

def test_parse_docx_bytes():
    # Generate a small in-memory valid DOCX
    doc = docx.Document()
    doc.add_paragraph("Jane Smith - Cloud Architect")
    doc.add_paragraph("Specialized in Kubernetes, Terraform, and GCP.")
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    result = DocumentParser.parse_from_bytes(docx_bytes, "candidate_resume.docx")
    assert "Jane Smith - Cloud Architect" in result
    assert "Kubernetes" in result

@pytest.mark.asyncio
async def test_parse_from_url_invalid():
    # Invalid remote URL should return graceful error text instead of throwing
    result = await DocumentParser.parse_from_url("http://invalid-non-existent-domain-12345.com/resume.pdf")
    assert "Failed to download and parse resume" in result
